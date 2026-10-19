"""Step 8b - Manuscript tables (BMC Nephrology format) generated from the analysis outputs.

Main tables  -> outputs/tables/THKFRE_main_tables.docx
  Table 1  Baseline characteristics by 2-year outcome
  Table 2  Predictive performance of the published KFRE and TH-KFRE
  Table 3  KFRE linear predictor and baseline risks (published and TH-KFRE)
  Table 4  Classification at clinical risk thresholds
Additional file 1 -> outputs/tables/THKFRE_additional_file_1.docx
  Table S1 Model updating: apparent and optimism-corrected performance, LR tests
  Table S2 Subgroup performance
  Table S3 Sensitivity analyses (competing risk handling; data corrections)
  Table S4 Temporal validation (2021 -> 2022)

BMC rules applied: Word table objects, no colour/shading, title above (<= 15 words),
legend/abbreviations below, no commas as thousands separators.
"""
import json

import numpy as np
import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Mm

from thkfre.config import (DERIVED_DIR, TABLE_DIR, HORIZON_DAYS, KFRE_BETA, KFRE_MEANS,
                           KFRE_S0_2Y_NORTH_AMERICA, KFRE_S0_2Y_NON_NORTH_AMERICA,
                           THKFRE_SLOPE, THKFRE_S0_2Y)
from thkfre.kfre import kfre_risk, thkfre_risk

H = HORIZON_DAYS
FONT = "Arial"


# ------------------------------------------------------------------ formatting helpers
def n_pct(mask, denom=None):
    m = np.asarray(mask, bool)
    d = len(m) if denom is None else denom
    return f"{m.sum()} ({100 * m.sum() / d:.1f})"


def mean_sd(x, dp=1):
    x = pd.Series(x).dropna()
    return f"{x.mean():.{dp}f} ± {x.std():.{dp}f}"


def med_iqr(x, dp=1):
    x = pd.Series(x).dropna()
    q1, q2, q3 = np.percentile(x, [25, 50, 75])
    return f"{q2:.{dp}f} ({q1:.{dp}f}–{q3:.{dp}f})"


def ci(v, lo, hi, dp=3):
    return f"{v:.{dp}f} ({lo:.{dp}f}–{hi:.{dp}f})"


def _border(cell, edges):
    tcPr = cell._tc.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for e in ("top", "bottom", "left", "right"):
        el = OxmlElement(f"w:{e}")
        if e in edges:
            el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "6"); el.set(qn("w:color"), "000000")
        else:
            el.set(qn("w:val"), "nil")
        b.append(el)
    tcPr.append(b)


def add_table(doc, title, header, rows, legend, col_widths_mm, bold_rows=(), font_pt=8.5):
    p = doc.add_paragraph()
    r = p.add_run(title); r.bold = True; r.font.size = Pt(10); r.font.name = FONT
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    layout = OxmlElement("w:tblLayout"); layout.set(qn("w:type"), "fixed")
    t._tbl.tblPr.append(layout)
    for j, w in enumerate(col_widths_mm):
        t.columns[j].width = Mm(w)
    for i, row in enumerate([header] + rows):
        for j, val in enumerate(row):
            c = t.cell(i, j)
            c.width = Mm(col_widths_mm[j])
            c.text = ""
            para = c.paragraphs[0]
            para.paragraph_format.space_after = Pt(1); para.paragraph_format.space_before = Pt(1)
            para.paragraph_format.line_spacing = 1.0
            run = para.add_run(str(val))
            run.font.size = Pt(font_pt); run.font.name = FONT
            if i == 0 or (i - 1) in bold_rows:
                run.bold = True
            edges = []
            if i == 0:
                edges = ["top", "bottom"]
            if i == len(rows):
                edges.append("bottom")
            _border(c, edges)
    lp = doc.add_paragraph()
    lr = lp.add_run(legend); lr.font.size = Pt(8); lr.font.name = FONT
    doc.add_paragraph()


def new_doc(landscape=False):
    doc = Document()
    sec = doc.sections[0]
    if landscape:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = Mm(297), Mm(210)
    else:
        sec.page_width, sec.page_height = Mm(210), Mm(297)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Mm(18))
    st = doc.styles["Normal"]; st.font.name = FONT; st.font.size = Pt(9)
    return doc


# ------------------------------------------------------------------ Table 1
def table1(doc, df):
    groups = [("All patients", df.index == df.index), ("Kidney failure", df.status == 1),
              ("Death before kidney failure", df.status == 2), ("Alive without kidney failure", df.status == 0)]
    header = ["Characteristic"] + [f"{g}\n(n = {int(m.sum())})" for g, m in groups]
    df = df.copy()
    df["p_nonna"] = kfre_risk(df, KFRE_S0_2Y_NON_NORTH_AMERICA) * 100
    df["p_th"] = thkfre_risk(df) * 100

    def row(label, fn):
        return [label] + [fn(df[m]) for _, m in groups]

    def stage(d, lo, hi):
        return n_pct((d.egfr >= lo) & (d.egfr < hi))

    rows = [
        row("Age, years, mean ± SD", lambda d: mean_sd(d.age)),
        row("Male sex, n (%)", lambda d: n_pct(d.male == 1)),
        row("eGFR, mL/min/1.73 m², mean ± SD", lambda d: mean_sd(d.egfr)),
        row("eGFR category, n (%)", lambda d: ""),
        row("   G3a (45–59)", lambda d: stage(d, 45, 60)),
        row("   G3b (30–44)", lambda d: stage(d, 30, 45)),
        row("   G4 (15–29)", lambda d: stage(d, 15, 30)),
        row("   G5 (<15)", lambda d: stage(d, 0, 15)),
        row("UACR, mg/g, median (IQR)", lambda d: med_iqr(d.uacr_mg_g)),
        row("Albuminuria category, n (%)", lambda d: ""),
        row("   A1 (<30 mg/g)", lambda d: n_pct(d.uacr_mg_g < 30)),
        row("   A2 (30–299 mg/g)", lambda d: n_pct((d.uacr_mg_g >= 30) & (d.uacr_mg_g < 300))),
        row("   A3 (≥300 mg/g)", lambda d: n_pct(d.uacr_mg_g >= 300)),
        row("Serum creatinine, mg/dL, median (IQR)", lambda d: med_iqr(d.creatinine, 2)),
    ]
    labs = [("Bicarbonate, mmol/L, mean ± SD", "bicarbonate", 1), ("Calcium, mg/dL, mean ± SD", "calcium", 1),
            ("Phosphate, mg/dL, mean ± SD", "phosphate", 1), ("Potassium, mmol/L, mean ± SD", "potassium", 2)]
    for lab, col, dp in labs:
        rows.append(row(f"{lab}ᵃ", lambda d, c=col, dp=dp: f"{mean_sd(d[c], dp)} [{d[c].notna().sum()}]"))
    rows += [
        row("Predicted 2-year risk, %, median (IQR)", lambda d: ""),
        row("   KFRE (non-North American)", lambda d: med_iqr(d.p_nonna)),
        row("   TH-KFRE", lambda d: med_iqr(d.p_th)),
        row("KRT modality, n (%)ᵇ", lambda d: ""),
        row("   Haemodialysis", lambda d: n_pct(d.krt_modality == "HD") if (d.status == 1).any() else "–"),
        row("   Peritoneal dialysis", lambda d: n_pct(d.krt_modality == "PD") if (d.status == 1).any() else "–"),
        row("Time to event or censoring, days, median (IQR)", lambda d: med_iqr(d.time_days, 0)),
    ]
    legend = ("Kidney failure: maintenance dialysis for more than 3 months or kidney transplantation within "
              "2 years. ᵃ[n] = number of patients with the measurement. ᵇPercentages of all patients in the column; "
              "no patient received a pre-emptive transplant. eGFR, estimated glomerular filtration rate (CKD-EPI); "
              "UACR, urine albumin-to-creatinine ratio; KRT, kidney replacement therapy; IQR, interquartile range; "
              "SD, standard deviation; KFRE, Kidney Failure Risk Equation.")
    add_table(doc, "Table 1 Baseline characteristics of the cohort by 2-year outcome", header, rows, legend,
              [58, 30, 30, 30, 30], bold_rows=())


# ------------------------------------------------------------------ Table 2
def table2(doc):
    v = pd.read_csv(TABLE_DIR / "02_validation_performance.csv").set_index("model")
    sens = pd.read_csv(TABLE_DIR / "06_sensitivity_data.csv")
    th = sens[(sens.subgroup == "Main analysis") & (sens.model == "TH-KFRE v1.0")].iloc[0]
    oc = pd.read_csv(TABLE_DIR / "04_optimism_corrected.csv").set_index("model").loc["M2 Recal. baseline+slope"]
    na, nn = v.loc["KFRE-NA (S0=0.9750)"], v.loc["KFRE-nonNA (S0=0.9832)"]
    header = ["Measure", "KFRE\n(North American)", "KFRE\n(non-North American)", "TH-KFRE\n(apparent)",
              "TH-KFRE\n(optimism-corrected)ᵃ"]
    rows = [
        ["Harrell's C", ci(na.Harrell_C, na.Harrell_C_lo, na.Harrell_C_hi), ci(nn.Harrell_C, nn.Harrell_C_lo, nn.Harrell_C_hi),
         ci(na.Harrell_C, na.Harrell_C_lo, na.Harrell_C_hi) + "ᵇ", f"{oc.corrected_Harrell_C:.3f}"],
        ["AUC at 2 years", ci(na.AUC_2y, na.AUC_2y_lo, na.AUC_2y_hi), ci(nn.AUC_2y, nn.AUC_2y_lo, nn.AUC_2y_hi),
         ci(na.AUC_2y, na.AUC_2y_lo, na.AUC_2y_hi) + "ᵇ", f"{oc.corrected_AUC_2y:.3f}"],
        ["Mean predicted risk, %", f"{100 * na.mean_predicted:.1f}", f"{100 * nn.mean_predicted:.1f}",
         f"{100 * th.mean_predicted:.1f}", "–"],
        ["Observed risk, %ᶜ", f"{100 * na.observed_AJ:.1f}", f"{100 * nn.observed_AJ:.1f}", f"{100 * th.observed_AJ:.1f}", "–"],
        ["O/E ratio", ci(na.O_E, na.O_E_lo, na.O_E_hi, 2), ci(nn.O_E, nn.O_E_lo, nn.O_E_hi, 2),
         ci(th.O_E, th.O_E_lo, th.O_E_hi, 2), f"{oc.corrected_O_E:.2f}"],
        ["Calibration slope", ci(na.cal_slope, na.cal_slope_lo, na.cal_slope_hi, 2),
         ci(nn.cal_slope, nn.cal_slope_lo, nn.cal_slope_hi, 2),
         ci(th.cal_slope, th.cal_slope_lo, th.cal_slope_hi, 2), f"{oc.corrected_cal_slope:.2f}"],
        ["ICI", f"{na.ICI:.3f}", f"{nn.ICI:.3f}", f"{oc.apparent_ICI:.3f}", f"{oc.corrected_ICI:.3f}"],
        ["Brier score", ci(na.Brier, na.Brier_lo, na.Brier_hi, 4), ci(nn.Brier, nn.Brier_lo, nn.Brier_hi, 4),
         f"{oc.apparent_Brier:.4f}", f"{oc.corrected_Brier:.4f}"],
        ["Scaled Brier score", ci(na.scaled_Brier, na.scaled_Brier_lo, na.scaled_Brier_hi, 3),
         ci(nn.scaled_Brier, nn.scaled_Brier_lo, nn.scaled_Brier_hi, 3), f"{oc.apparent_scaled_Brier:.3f}",
         f"{oc.corrected_scaled_Brier:.3f}"],
    ]
    legend = ("Values are estimates (95% bootstrap confidence intervals, 1000 resamples). Death before kidney failure "
              "was treated as a competing event. ᵃBootstrap optimism correction with the whole updating procedure "
              "repeated in each of 1000 resamples. ᵇTH-KFRE ranks patients identically to the KFRE (monotone "
              "transformation of the same linear predictor), so discrimination is identical. ᶜAalen–Johansen cumulative "
              "incidence of kidney failure at 2 years. AUC, time-dependent area under the ROC curve (inverse probability "
              "of censoring weighted); O/E, observed-to-expected ratio; ICI, integrated calibration index.")
    add_table(doc, "Table 2 Predictive performance of the published KFRE and TH-KFRE at 2 years",
              header, rows, legend, [36, 34, 34, 34, 32])


# ------------------------------------------------------------------ Table 3
def table3(doc):
    spec = json.loads((TABLE_DIR / "03_model_specifications.json").read_text())
    labels = {"age10": "Age, per 10 years", "male": "Male sex", "egfr5": "eGFR, per 5 mL/min/1.73 m²",
              "ln_uacr": "ln UACR (mg/g)"}
    header = ["Predictor", "Centring value", "KFRE coefficient (HR)", "TH-KFRE effective coefficient (HR)ᵃ"]
    rows = [[labels[v], f"{KFRE_MEANS[v]}", f"{KFRE_BETA[v]:.4f} ({np.exp(KFRE_BETA[v]):.2f})",
             f"{THKFRE_SLOPE * KFRE_BETA[v]:.4f} ({np.exp(THKFRE_SLOPE * KFRE_BETA[v]):.2f})"] for v in labels]
    rows += [["Baseline 2-year survival, S₀", "", f"{KFRE_S0_2Y_NORTH_AMERICA} (North American)\n"
              f"{KFRE_S0_2Y_NON_NORTH_AMERICA} (non-North American)", f"{THKFRE_S0_2Y}"],
             ["Slope on KFRE linear predictor", "", "1 (fixed)",
              f"{THKFRE_SLOPE} (SE {spec['M2']['slope_se']:.3f})"]]
    legend = ("KFRE linear predictor LP = −0.2201 × (age/10 − 7.036) + 0.2467 × (male − 0.5642) − 0.5567 × (eGFR/5 − 7.222) "
              "+ 0.4510 × (ln UACR − 5.137). KFRE 2-year risk = 1 − S₀^exp(LP). "
              f"TH-KFRE 2-year risk = 1 − {THKFRE_S0_2Y}^exp({THKFRE_SLOPE} × LP). "
              "ᵃKFRE coefficient multiplied by the TH-KFRE slope; shown for interpretation only. "
              "HR, hazard ratio; SE, standard error.")
    add_table(doc, "Table 3 KFRE linear predictor and TH-KFRE recalibration parameters",
              header, rows, legend, [46, 26, 50, 50])


# ------------------------------------------------------------------ Table 4
def table4(doc):
    t = pd.read_csv(TABLE_DIR / "05_threshold_performance.csv")
    t = t[t.threshold.isin([0.10, 0.20, 0.40])]
    names = {"KFRE-NA": "KFRE (North American)", "KFRE-nonNA": "KFRE (non-North American)", "TH-KFRE v1.0": "TH-KFRE"}
    header = ["Threshold / model", "Flagged, n (%)", "Sensitivity", "Specificity", "PPV", "NPV", "NNR"]
    rows, bold = [], []
    for c in [0.10, 0.20, 0.40]:
        bold.append(len(rows))
        rows.append([f"2-year risk ≥{c:.0%}", "", "", "", "", "", ""])
        for m in ["KFRE-NA", "KFRE-nonNA", "TH-KFRE v1.0"]:
            r = t[(t.threshold == c) & (t.model == m)].iloc[0]
            rows.append([f"   {names[m]}", f"{int(r.n_flagged)} ({100 * r.pct_flagged:.1f})",
                         ci(r.sensitivity, r.sensitivity_lo, r.sensitivity_hi, 2),
                         ci(r.specificity, r.specificity_lo, r.specificity_hi, 2),
                         ci(r.PPV, r.PPV_lo, r.PPV_hi, 2), ci(r.NPV, r.NPV_lo, r.NPV_hi, 2),
                         ci(r.NNR, r.NNR_lo, r.NNR_hi, 1)])
    legend = ("Values are estimates (95% bootstrap confidence intervals). 10%: referral to multidisciplinary CKD care; "
              "40%: preparation for kidney replacement therapy. Sensitivity and specificity are inverse probability of "
              "censoring weighted; PPV and NPV are Aalen–Johansen cumulative incidences in the flagged and unflagged groups. "
              "PPV, positive predictive value; NPV, negative predictive value; NNR, number needed to refer (1/PPV).")
    add_table(doc, "Table 4 Classification performance at clinical 2-year risk thresholds",
              header, rows, legend, [40, 20, 23, 23, 23, 23, 22], bold_rows=bold, font_pt=8)


# ------------------------------------------------------------------ Additional file tables
def supplementary(doc):
    app = pd.read_csv(TABLE_DIR / "04_optimism_corrected.csv")
    lr = pd.read_csv(TABLE_DIR / "03_lr_tests.csv")
    names = {"M0a KFRE-NA": "M0a KFRE (North American)", "M0b KFRE-nonNA": "M0b KFRE (non-North American)",
             "M1 Recal. baseline": "M1 Recalibrated baseline", "M2 Recal. baseline+slope": "M2 Baseline + slope (TH-KFRE)",
             "M3a Refit Cox": "M3a Refit, cause-specific Cox", "M3b Refit Fine-Gray": "M3b Refit, Fine–Gray"}
    header = ["Model", "C (app / corr)", "O/E (app / corr)", "Slope (app / corr)", "ICI (app / corr)", "Brier (app / corr)"]
    rows = [[names[r.model], f"{r.apparent_Harrell_C:.3f} / {r.corrected_Harrell_C:.3f}",
             f"{r.apparent_O_E:.2f} / {r.corrected_O_E:.2f}", f"{r.apparent_cal_slope:.2f} / {r.corrected_cal_slope:.2f}",
             f"{r.apparent_ICI:.3f} / {r.corrected_ICI:.3f}", f"{r.apparent_Brier:.4f} / {r.corrected_Brier:.4f}"]
            for r in app.itertuples()]
    legend = ("app, apparent (full cohort); corr, bootstrap optimism-corrected (1000 resamples). Likelihood-ratio tests "
              "(Fine–Gray partial likelihood): " + "; ".join(f"{r.test}: χ² = {r.LR_chi2:.2f}, df = {r.df}, p = {r.p:.4f}"
                                                           for r in lr.itertuples()) + ".")
    add_table(doc, "Table S1 Model updating: apparent and optimism-corrected performance", header, rows, legend,
              [52, 26, 24, 24, 26, 30])

    sub = pd.read_csv(TABLE_DIR / "06_subgroups.csv")
    header = ["Subgroup / model", "n", "KF events", "Harrell's C", "O/E (95% CI)", "Slope (95% CI)"]
    rows, bold = [], []
    for g in sub.subgroup.unique():
        bold.append(len(rows))
        rows.append([g.replace(">=", "≥"), "", "", "", "", ""])
        short = {"KFRE-NA": "KFRE (North American)", "KFRE-nonNA": "KFRE (non-North American)", "TH-KFRE v1.0": "TH-KFRE"}
        for r in sub[sub.subgroup == g].itertuples():
            rows.append([f"   {short[r.model]}", f"{r.n}", f"{r.KF_events}",
                         f"{r.Harrell_C:.3f}", ci(r.O_E, r.O_E_lo, r.O_E_hi, 2),
                         ci(r.cal_slope, r.cal_slope_lo, r.cal_slope_hi, 2) if pd.notna(r.cal_slope) else "n/a"])
    add_table(doc, "Table S2 Performance in pre-specified subgroups", header, rows,
              "Age 65–74 and ≥75 years were added post hoc. Calibration slopes not estimated with fewer than 20 events (n/a). "
              "TH-KFRE results are apparent (derived in the full cohort).", [56, 14, 18, 22, 32, 32], bold_rows=bold)

    cr = pd.read_csv(TABLE_DIR / "06_sensitivity_competing_risk.csv")
    header = ["Group", "n", "Deaths", "Observed risk, AJ (%)", "Observed risk, 1−KM (%)", "TH-KFRE O/E: AJ / KM"]
    rows = [[r["group"], f"{r['n']}", f"{r['deaths']}", f"{100 * r['observed_AJ']:.2f}", f"{100 * r['observed_KM']:.2f}",
             f"{r['O_E_AJ TH-KFRE v1.0']:.2f} / {r['O_E_KM TH-KFRE v1.0']:.2f}"] for _, r in cr.iterrows()]
    add_table(doc, "Table S3 Sensitivity analysis: death as competing event versus censoring", header, rows,
              "AJ, Aalen–Johansen (death as competing event); KM, Kaplan–Meier (death censored).",
              [30, 16, 18, 34, 34, 36])

    tv = pd.read_csv(TABLE_DIR / "04_temporal_validation.csv")
    header = ["Model (developed 2021, tested 2022)", "Harrell's C", "O/E (95% CI)", "Slope", "ICI", "Brier"]
    rows = [[names.get(r.model, r.model.replace("TH-KFRE v1.0", "TH-KFRE")), f"{r.Harrell_C:.3f}", ci(r.O_E, r.O_E_lo, r.O_E_hi, 2),
             f"{r.cal_slope:.2f}", f"{r.ICI:.3f}", f"{r.Brier:.4f}"] for r in tv.itertuples()]
    add_table(doc, "Table S4 Temporal validation: models developed in 2021, tested in 2022", header, rows,
              "Development: 2021 index visits (n = 389, 37 events); test: 2022 index visits (n = 668, 53 events).",
              [62, 22, 32, 18, 18, 20])


def main():
    df = pd.read_csv(DERIVED_DIR / "analysis_cohort.csv")
    doc = new_doc()
    table1(doc, df)
    doc.add_page_break()
    table2(doc)
    doc.add_page_break()
    table3(doc)
    doc.add_page_break()
    table4(doc)
    doc.save(TABLE_DIR / "THKFRE_main_tables.docx")
    sup = new_doc()
    supplementary(sup)
    sup.save(TABLE_DIR / "THKFRE_additional_file_1.docx")
    print("saved")


if __name__ == "__main__":
    main()
